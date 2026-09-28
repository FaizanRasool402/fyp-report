"""Insert methodology diagrams into the FYP PowerPoint."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import shutil
import os

SRC = r"c:\Users\Dell\Desktop\fyp-report\presentation\AI_Lung_Disease_Detection.pptx"
OUT = r"c:\Users\Dell\Desktop\fyp-report\presentation\AI_Lung_Disease_Detection_with_diagrams.pptx"
DIAG = r"c:\Users\Dell\Desktop\fyp-report\diagrams"

shutil.copy2(SRC, OUT)
prs = Presentation(OUT)


def replace_first_picture(slide, image_path):
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            el = shape._element
            el.getparent().remove(el)
            slide.shapes.add_picture(image_path, left, top, width=width, height=height)
            return True
    return False


def add_footer(slide, text="AI Agent For Lung Diseases Detection"):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip() == text:
            return
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Emu(6386286), prs.slide_width, Emu(471714)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(32, 56, 78)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER


def make_diagram_slide(prs, title, image_path, insert_after_index):
    blank = prs.slide_layouts[-1]
    for layout in prs.slide_layouts:
        if "blank" in layout.name.lower():
            blank = layout
            break

    slide = prs.slides.add_slide(blank)

    title_box = slide.shapes.add_textbox(0, Emu(18255), prs.slide_width, Emu(1100000))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(32, 56, 78)
    p.alignment = PP_ALIGN.CENTER

    left = Emu(500000)
    top = Emu(1400000)
    max_w = prs.slide_width - Emu(1000000)
    max_h = Emu(4700000)
    pic = slide.shapes.add_picture(image_path, left, top, width=max_w)
    if pic.height > max_h:
        ratio = float(max_h) / float(pic.height)
        pic.width = int(pic.width * ratio)
        pic.height = int(max_h)
    pic.left = int((prs.slide_width - pic.width) / 2)

    add_footer(slide)

    sldIdLst = prs.slides._sldIdLst
    sldIds = list(sldIdLst)
    new_sld = sldIds[-1]
    sldIdLst.remove(new_sld)
    target = insert_after_index + 1
    sldIdLst.insert(target, new_sld)
    return slide


# Replace existing pictures with updated diagrams
replacements = {
    15: os.path.join(DIAG, "06_system_inference.png"),
    18: os.path.join(DIAG, "03_model_architecture.png"),
    31: os.path.join(DIAG, "07_xray_gate.png"),
}
for idx, path in replacements.items():
    ok = replace_first_picture(prs.slides[idx], path)
    print(f"Replace slide {idx + 1}: {ok} <- {os.path.basename(path)}")

# Add pipeline diagram onto Methodology slide
meth = prs.slides[16]
has_pic = any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in meth.shapes)
if not has_pic:
    for shape in meth.shapes:
        if shape.has_text_frame and "shared dataset" in shape.text_frame.text.lower():
            shape.top = Emu(1250000)
            shape.height = Emu(1200000)
            break
    pic = meth.shapes.add_picture(
        os.path.join(DIAG, "01_methodology_pipeline.png"),
        Emu(600000),
        Emu(2600000),
        width=Emu(11000000),
    )
    if pic.height > Emu(3500000):
        ratio = float(Emu(3500000)) / float(pic.height)
        pic.width = int(pic.width * ratio)
        pic.height = int(Emu(3500000))
    pic.left = int((prs.slide_width - pic.width) / 2)
    print("Added pipeline to Methodology slide 17")

# Insert from highest index first so earlier indices stay stable
inserts = [
    (35, "Evaluation & Leakage-Probe Flow", "05_evaluation_flow.png"),
    (30, "End-to-End Application Inference", "06_system_inference.png"),
    (20, "Training Protocol (Two Phases)", "04_training_phases.png"),
    (11, "Dataset Unification Diagram", "02_data_unification.png"),
]

for after_idx, title, fname in inserts:
    path = os.path.join(DIAG, fname)
    make_diagram_slide(prs, title, path, after_idx)
    print(f"Inserted after slide {after_idx + 1}: {title}")

prs.save(OUT)
print("Saved", OUT)
print("Total slides", len(prs.slides))

for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    texts.append(p.text.strip())
                    break
        if texts:
            break
    title = texts[0] if texts else "(no text)"
    if 10 <= i <= 25 or 30 <= i <= 42:
        print(f"{i:02d}: {title}")
