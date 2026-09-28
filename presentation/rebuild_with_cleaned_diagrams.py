"""Rebuild presentation from original with cleaned diagrams (aspect preserved)."""
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
import shutil
import os

ORIG = r"c:\Users\Dell\Downloads\AI Lung Disease Detection - Final (revised) (5) (1) (3).pptx"
OUT = r"c:\Users\Dell\Desktop\fyp-report\presentation\AI_Lung_Disease_Detection_with_diagrams.pptx"
DESK = r"c:\Users\Dell\Desktop\AI_Lung_Disease_Detection_with_diagrams.pptx"
CLEAN = r"c:\Users\Dell\Desktop\fyp-report\diagrams\cleaned"

# 0-based slide index -> cleaned image
REPLACEMENTS = {
    14: "preprocessing_pipeline.png",       # Preprocessing Pipeline
    15: "system_architecture.png",          # Overall System Architecture
    18: "model_architectures.png",          # Methodology: Model Architectures
    21: "efficientnetb3_method.png",        # Methodology: EfficientNetB3
    25: "densenet121_method.png",           # Methodology: DenseNet121
    27: "custom_cnn_method.png",            # Methodology: Custom CNN
    31: "xray_gate_pipeline.png",           # Input Validation Pipeline
}


def fit_picture(slide, image_path, box_left, box_top, box_width, box_height):
    """Place image centered in box without stretching."""
    with Image.open(image_path) as im:
        iw, ih = im.size
    box_aspect = box_width / box_height
    img_aspect = iw / ih
    if img_aspect > box_aspect:
        w = box_width
        h = int(box_width / img_aspect)
    else:
        h = box_height
        w = int(box_height * img_aspect)
    left = int(box_left + (box_width - w) / 2)
    top = int(box_top + (box_height - h) / 2)
    slide.shapes.add_picture(image_path, left, top, width=w, height=h)


shutil.copy2(ORIG, OUT)
prs = Presentation(OUT)

for idx, fname in REPLACEMENTS.items():
    path = os.path.join(CLEAN, fname)
    slide = prs.slides[idx]
    replaced = False
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            el = shape._element
            el.getparent().remove(el)
            fit_picture(slide, path, left, top, width, height)
            replaced = True
            break
    print(f"slide {idx + 1}: {fname} -> {replaced}")

prs.save(OUT)
shutil.copy2(OUT, DESK)
print("saved", OUT)
print("desktop", DESK)
print("slides", len(prs.slides))
